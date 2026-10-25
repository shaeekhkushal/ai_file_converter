from pptx import Presentation

def process_ppt(file_path, max_records=None):
    """Extract text from each slide of a PowerPoint file.

    Returns a list of dictionaries with ``slide`` (1‑based index) and ``text``
    (concatenated text of all shapes on that slide). Empty slides are omitted.
    """
    prs = Presentation(file_path)
    data = []
    for i, slide in enumerate(prs.slides, start=1):
        if max_records is not None and i > max_records:
            break
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                texts.append(shape.text.strip())
        if texts:
            data.append({"slide": i, "text": "\n".join(texts)})
    return data
